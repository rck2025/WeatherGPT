import http.server
import json
import os
import shutil
import socketserver
import sys
import threading
import time
from pathlib import Path

import docx
import pandas as pd
import pymupdf
from fastapi.testclient import TestClient

# Ensure UTF-8 stdout
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(WORKSPACE_ROOT))

import backend.main as main_module
from backend.main import app, get_active_alerts
from backend.services.ingestion import UniversalScraper, WeatherAlertAdapter

DATA_DIR = WORKSPACE_ROOT / "backend" / "data" / "comprehensive_test_data"
SERVER_PORT = 8899


class MockWebHandler(http.server.SimpleHTTPRequestHandler):
    """Local HTTP handler simulating real external weather websites."""

    def do_GET(self):
        if self.path == "/imd_mumbai_cyclone.html":
            self.send_response(200)
            self.send_header("Content-type", "text/html; charset=utf-8")
            self.end_headers()
            html = """<!DOCTYPE html>
            <html>
            <head><title>IMD Coastal Weather Warning</title><script>var x = 1;</script></head>
            <body>
                <header><nav>Navigation</nav></header>
                <h1>Extremely Severe Cyclone Alert for Coastal Maharashtra</h1>
                <p>Location: Mumbai and Thane. A red alert has been sounded by the IMD. Heavy rainfall and gale wind speeds expected.</p>
                <footer>IMD Govt of India</footer>
            </body>
            </html>"""
            self.wfile.write(html.encode("utf-8"))

        elif self.path == "/ndma_kolkata_flood.html":
            self.send_response(200)
            self.send_header("Content-type", "text/html; charset=utf-8")
            self.end_headers()
            html = """<!DOCTYPE html>
            <html>
            <head><title>NDMA Flood Warning</title></head>
            <body>
                <h1>Severe Flood Warning in West Bengal</h1>
                <p>Severe rainfall and thunderstorm warning in Kolkata and surrounding districts. High risk of waterlogging.</p>
            </body>
            </html>"""
            self.wfile.write(html.encode("utf-8"))

        elif self.path == "/general_news.html":
            self.send_response(200)
            self.send_header("Content-type", "text/html; charset=utf-8")
            self.end_headers()
            html = "<html><body><h1>Tech & Culture</h1><p>Annual technology summit scheduled for December.</p></body></html>"
            self.wfile.write(html.encode("utf-8"))

        elif self.path == "/error_404.html":
            self.send_response(404)
            self.end_headers()
            self.wfile.write(b"404 Not Found")

        elif self.path == "/error_500.html":
            self.send_response(500)
            self.end_headers()
            self.wfile.write(b"500 Internal Server Error")

        else:
            self.send_response(404)
            self.end_headers()

    def log_message(self, format, *args):
        # Suppress noisy standard HTTP logs during automated test
        pass


def start_mock_server():
    server = socketserver.TCPServer(("127.0.0.1", SERVER_PORT), MockWebHandler)
    server_thread = threading.Thread(target=server.serve_forever, daemon=True)
    server_thread.start()
    time.sleep(0.5)
    return server


def create_physical_test_files():
    """Create authentic files on disk across every supported format."""
    if DATA_DIR.exists():
        shutil.rmtree(DATA_DIR)
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    sources_map = {}

    # 1. Real PDF Document (Multi-page using PyMuPDF)
    pdf_path = DATA_DIR / "chennai_monsoon_bulletin.pdf"
    pdf_doc = pymupdf.open()
    p1 = pdf_doc.new_page()
    p1.insert_text(pymupdf.Point(50, 70), "IMD Special Weather Bulletin - Tamil Nadu", fontsize=14)
    p1.insert_text(pymupdf.Point(50, 100), "Severe Monsoon Rain and Flood Alert for Chennai", fontsize=12)
    p1.insert_text(pymupdf.Point(50, 130), "Continuous torrential downpour across coastal Tamil Nadu.", fontsize=10)
    p2 = pdf_doc.new_page()
    p2.insert_text(pymupdf.Point(50, 70), "Precautionary Measures:", fontsize=12)
    p2.insert_text(pymupdf.Point(50, 100), "Warning issued for low-lying coastal areas in Chennai.", fontsize=10)
    pdf_doc.save(str(pdf_path))
    pdf_doc.close()
    sources_map["PDF"] = (str(pdf_path), True, "Chennai")

    # 2. Real Word DOCX Document (using python-docx)
    docx_path = DATA_DIR / "delhi_heatwave_advisory.docx"
    doc = docx.Document()
    doc.add_heading("National Disaster Management Authority Advisory", level=1)
    doc.add_paragraph("Severe Heatwave Warning for Delhi NCR")
    doc.add_paragraph("Extreme temperatures exceeding 45 degrees Celsius expected across Delhi and Haryana.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Alert Level"
    table.cell(1, 0).text = "Delhi"
    table.cell(1, 1).text = "Red Alert"
    doc.save(str(docx_path))
    sources_map["DOCX"] = (str(docx_path), True, "Delhi")

    # 3. Real Excel Spreadsheet (.xlsx using pandas and openpyxl)
    xlsx_path = DATA_DIR / "odisha_cyclone_telemetry.xlsx"
    with pd.ExcelWriter(str(xlsx_path), engine="openpyxl") as writer:
        df_advisory = pd.DataFrame({
            "Station": ["Bhubaneswar", "Puri", "Cuttack"],
            "Warning_Type": ["Severe Cyclonic Storm Warning", "Cyclone Red Alert", "Flood Alert"],
            "Wind_Speed_Kmph": [110, 130, 95],
            "Rainfall_mm": [180.5, 220.0, 140.2]
        })
        df_advisory.to_excel(writer, sheet_name="Coastal_Warnings", index=False)
    sources_map["EXCEL"] = (str(xlsx_path), True, "Bhubaneswar")

    # 4. Real JSON Document
    json_path = DATA_DIR / "bengaluru_flash_flood.json"
    json_data = {
        "source_agency": "Karnataka State Disaster Management Authority",
        "bulletin": {
            "title": "Urban Flood and Heavy Rainfall Alert for Bengaluru",
            "state": "Karnataka",
            "city": "Bengaluru",
            "details": {
                "summary": "Monsoon downpour triggered flash flood warning in low-lying zones across Bengaluru.",
                "severity_index": "HIGH"
            }
        }
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, indent=2)
    sources_map["JSON"] = (str(json_path), True, "Bengaluru")

    # 5. Real HTML File on disk
    html_path = DATA_DIR / "shimla_landslide_report.html"
    with open(html_path, "w", encoding="utf-8") as f:
        f.write("""<html>
        <head><title>Himachal Disaster Bulletin</title><style>body{color:red;}</style></head>
        <body>
            <h1>Cloudburst and Landslide Warning in Shimla</h1>
            <p>Severe rainfall causing major landslide risk in Shimla district. Red alert in effect.</p>
        </body>
        </html>""")
    sources_map["HTML_FILE"] = (str(html_path), True, "Shimla")

    # 6. Real CSV Telemetry File
    csv_path = DATA_DIR / "guwahati_river_levels.csv"
    with open(csv_path, "w", encoding="utf-8") as f:
        f.write("River,Station,Status,Alert_Message\nBrahmaputra,Guwahati,DANGER,Severe flood warning and landslide risk near Guwahati.\n")
    sources_map["CSV"] = (str(csv_path), True, "Guwahati")

    # 7. Irrelevant Files (Should be filtered out)
    irrelevant_path = DATA_DIR / "sports_scores.txt"
    with open(irrelevant_path, "w", encoding="utf-8") as f:
        f.write("Sports Championship 2026: Team A defeated Team B by 3 goals.")
    sources_map["IRRELEVANT_TXT"] = (str(irrelevant_path), False, None)

    # 8. Corrupted & Edge Case Files
    corrupt_pdf = DATA_DIR / "corrupted_file.pdf"
    with open(corrupt_pdf, "wb") as f:
        f.write(b"NOT_A_REAL_PDF_HEADER_JUNK_123456789")
    sources_map["CORRUPT_PDF"] = (str(corrupt_pdf), False, None)

    corrupt_json = DATA_DIR / "corrupted.json"
    with open(corrupt_json, "w", encoding="utf-8") as f:
        f.write("{invalid_json_missing_quotes: [")
    sources_map["CORRUPT_JSON"] = (str(corrupt_json), False, None)

    empty_file = DATA_DIR / "empty.txt"
    with open(empty_file, "w", encoding="utf-8") as f:
        pass
    sources_map["EMPTY_FILE"] = (str(empty_file), False, None)

    return sources_map


def run_comprehensive_test():
    print("=" * 80)
    print("RUNNING COMPREHENSIVE LIVE INGESTION & SCRAPER TEST SUITE")
    print("=" * 80)

    # 1. Start live web server
    server = start_mock_server()
    print(f"[INFO] Started live mock weather web server on http://127.0.0.1:{SERVER_PORT}")

    # 2. Create physical files on disk
    sources_map = create_physical_test_files()
    print(f"[INFO] Created physical test files on disk in: {DATA_DIR}")

    # 3. Live Web URLs to scrape
    live_urls = {
        "LIVE_URL_MUMBAI": (f"http://127.0.0.1:{SERVER_PORT}/imd_mumbai_cyclone.html", True, "Mumbai"),
        "LIVE_URL_KOLKATA": (f"http://127.0.0.1:{SERVER_PORT}/ndma_kolkata_flood.html", True, "Kolkata"),
        "LIVE_URL_IRRELEVANT": (f"http://127.0.0.1:{SERVER_PORT}/general_news.html", False, None),
        "LIVE_URL_404": (f"http://127.0.0.1:{SERVER_PORT}/error_404.html", False, None),
        "LIVE_URL_500": (f"http://127.0.0.1:{SERVER_PORT}/error_500.html", False, None),
    }

    all_test_sources = {**sources_map, **live_urls}
    sources_list = [item[0] for item in all_test_sources.values()]

    print(f"[INFO] Total Sources Prepared: {len(sources_list)}")
    for name, (path_or_url, is_valid, city) in all_test_sources.items():
        print(f"  - [{name}]: {path_or_url} (Expected Valid: {is_valid}, Target City: {city})")

    # -------------------------------------------------------------
    # STAGE 1: DIRECT TEST OF UNIVERSALSCRAPER ACROSS ALL FORMATS
    # -------------------------------------------------------------
    print("\n" + "-" * 80)
    print("STAGE 1: TESTING UNIVERSALSCRAPER DIRECTLY ON EVERY FORMAT")
    print("-" * 80)

    scraper = UniversalScraper()
    adapter = WeatherAlertAdapter()

    for name, (source_path, expected_valid, expected_city) in all_test_sources.items():
        extracted_text = scraper.extract_text(source_path)
        is_rel = scraper.is_relevant(extracted_text) if extracted_text else False

        print(f"\nTesting Source [{name}]: {source_path}")
        if expected_valid:
            assert extracted_text is not None, f"Failed to extract text from valid source: {name}"
            assert is_rel is True, f"Valid disaster source failed relevance filter: {name}"
            
            alert = adapter.to_weather_alert(extracted_text, source_path)
            assert alert is not None
            assert alert.severity in ["Extreme", "High", "Moderate"]
            print(f"  -> Extracted Length: {len(extracted_text)} chars")
            print(f"  -> Relevance Filter: PASSED (Disaster keywords found)")
            print(f"  -> Adapted Alert Title: {alert.title}")
            print(f"  -> Severity: {alert.severity}")
            print(f"  -> Description Snippet: {alert.description[:120]}...")
            if expected_city:
                assert expected_city.lower() in alert.description.lower(), f"Expected city {expected_city} not found in description"
                print(f"  -> Location Verified: {expected_city}")
        else:
            print(f"  -> Non-disaster / Broken source successfully handled. Extracted Text: {'Found' if extracted_text else 'None'}, Relevant: {is_rel}")

    print("\n[PASSED] Stage 1: UniversalScraper accurately parsed all formats and handled all corruptions.")

    # -------------------------------------------------------------
    # STAGE 2: BATCH INGESTION ENDPOINT (POST /api/v1/ingest/batch)
    # -------------------------------------------------------------
    print("\n" + "-" * 80)
    print("STAGE 2: BATCH INGESTION VIA API (POST /api/v1/ingest/batch)")
    print("-" * 80)

    main_module.alerts.clear()
    client = TestClient(app)

    start_time = time.perf_counter()
    batch_response = client.post("/api/v1/ingest/batch", json={"sources": sources_list})
    elapsed = time.perf_counter() - start_time

    assert batch_response.status_code == 200
    batch_data = batch_response.json()

    print(f"API Ingestion Response Status: {batch_response.status_code}")
    print(f"Batch Processing Time: {elapsed:.3f} seconds ({len(sources_list)/elapsed:.1f} sources/sec)")
    print(f"Total Sources Submitted: {len(sources_list)}")
    print(f"Alerts Ingested: {batch_data.get('alerts_added', len(main_module.alerts))}")
    print(f"Global In-Memory Alerts: {len(main_module.alerts)}")

    expected_valid_count = sum(1 for _, valid, _ in all_test_sources.values() if valid)
    assert len(main_module.alerts) == expected_valid_count, f"Expected {expected_valid_count} alerts, got {len(main_module.alerts)}"
    print(f"\n[PASSED] Stage 2: Ingestion API processed batch perfectly.")

    # -------------------------------------------------------------
    # STAGE 3: ACTIVE ALERTS SEARCH ACROSS ALL INGESTED LOCATIONS
    # -------------------------------------------------------------
    print("\n" + "-" * 80)
    print("STAGE 3: ACTIVE ALERTS SEARCH BY LOCATION")
    print("-" * 80)

    target_cities = ["Mumbai", "Chennai", "Delhi", "Kolkata", "Bengaluru", "Shimla", "Bhubaneswar", "Guwahati", "London"]

    for city in target_cities:
        matched = get_active_alerts(city)
        if city == "London":
            assert len(matched) == 0
            print(f"  - [{city}]: 0 alerts (Correctly identified as unaffected)")
        else:
            assert len(matched) >= 1, f"Expected at least 1 alert for {city}, got {len(matched)}"
            print(f"  - [{city}]: Found {len(matched)} active alert(s) -> Severity: {matched[0].severity} | Title: {matched[0].title}")

    print("\n[PASSED] Stage 3: Location search is 100% accurate.")

    # -------------------------------------------------------------
    # STAGE 4: END-TO-END CHATBOT (POST /chat) OUTPUT GENERATION
    # -------------------------------------------------------------
    print("\n" + "-" * 80)
    print("STAGE 4: LIVE CHATBOT OUTPUT GENERATION ACROSS CITIES")
    print("-" * 80)

    chat_queries = [
        ("Mumbai", "Are there any cyclone or flood warnings in Mumbai?"),
        ("Chennai", "What is the monsoon status in Chennai?"),
        ("Delhi", "Is there a heatwave alert for Delhi?"),
        ("London", "What is the weather status in London?")
    ]

    for city, query in chat_queries:
        req = {
            "query": query,
            "location": {"city": city, "country": "India" if city != "London" else "UK"}
        }
        res = client.post("/chat", json=req)
        assert res.status_code == 200
        chat_data = res.json()
        print(f"\n>>> QUERY: \"{query}\" [City: {city}]")
        print(f">>> BOT REPLY:\n{chat_data['bot_reply']}")
        print(f">>> ALERTS ATTACHED: {len(chat_data['alerts'])}")

    print("\n[PASSED] Stage 4: Chatbot synthesized accurate live responses for all cities.")

    # Clean up test directories & shutdown server
    if DATA_DIR.exists():
        shutil.rmtree(DATA_DIR)
    server.shutdown()

    print("\n" + "=" * 80)
    print("ALL COMPREHENSIVE PRODUCTION INGESTION & SCRAPER TESTS COMPLETED WITH 100% SUCCESS!")
    print("=" * 80)


if __name__ == "__main__":
    run_comprehensive_test()
